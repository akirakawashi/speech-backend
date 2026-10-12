"""Файл «Сверка оценок»: что решила модель по этапу и что при этом было в разговоре.

Файл — для контролёра. Он читает строку: этап, оценка модели, её пояснение и
реплики, на которые она сослалась, — и ставит рядом свою оценку. По
расхождениям видно, каким этапам у модели можно верить, а каким нет.

Листов три. «Сверка» — строка на пару «звонок × этап», с двумя пустыми
колонками для контролёра. «Разговоры» — полный текст тех же звонков: «этап не
выполнен» значит «этого не сказали», и по одним цитатам отсутствие не
проверить. «Этапы» — что требуется по каждому этапу: это то же требование,
которое получила модель, и без него контролёр сверял бы оценку со своим
пониманием этапа, а не с тем, что проверялось.

Какие звонки попадут в файл, решает вызывающий (`quality_routes`): отбор у
него общий со сводкой. Здесь — только реплики и сборка.

В файл идёт текст расшифровки как есть — тот же, что видела модель.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Mapping, Sequence
from datetime import datetime
from typing import Any, Final
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from src.application.dto.quality import CallScript
from src.infrastructure.reports.calls import SPEAKER_LABEL, fit_text

#: Сколько звонков в файле по умолчанию и не больше скольких. Тридцать звонков
#: — триста строк сверки: объём, который контролёр проходит за один раз. Сотня —
#: потолок: файл собирается в памяти и целиком.
REVIEW_SIZE: Final[int] = 30
REVIEW_MAX: Final[int] = 100

NOT_NEEDED: Final[str] = "не понадобился"

CHECK_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("№", 5),
    ("Звонок", 17),
    ("Сотрудник", 24),
    ("Длительность", 13),
    ("Запись и разбор", 16),
    ("Этап скрипта", 38),
    ("Обязательный", 13),
    ("Оценка модели", 15),
    ("Пояснение модели", 48),
    ("Реплики, на которые сослалась модель", 70),
    ("Оценка контролёра (1 / 0 / не понадобился)", 20),
    ("Комментарий контролёра", 40),
)
#: Колонки, которые заполняет контролёр: шапка у них выделена цветом.
_CONTROLLER_COLUMNS: Final[frozenset[int]] = frozenset({11, 12})

TALK_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("№", 5),
    ("Звонок", 17),
    ("Сотрудник", 24),
    ("Время", 8),
    ("Кто говорит", 14),
    ("Реплика", 110),
)

STAGE_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("Этап скрипта", 38),
    ("Обязательный", 13),
    ("Что требуется", 80),
    ("Когда этап нужен", 55),
)

_MOMENT_FORMAT: Final[str] = "DD.MM.YYYY HH:MM"
#: Звонки чередуются фоном: десять строк одного звонка читаются одним блоком.
_BAND: Final[PatternFill] = PatternFill("solid", fgColor="F3F4F6")
_TO_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFF2CC")
_TOP: Final[Alignment] = Alignment(vertical="top")
_TOP_WRAPPED: Final[Alignment] = Alignment(vertical="top", wrap_text=True)

_LINES_SQL = text(
    """
    SELECT call_id, seq, speaker, start_ms, text_raw
      FROM speech_transcript_segments
     WHERE call_id = ANY(CAST(:call_ids AS uuid[]))
     ORDER BY call_id, seq
    """
)


class ReviewLine(BaseModel):
    """Реплика и таймкод для файла ручной сверки оценки."""

    model_config = ConfigDict(frozen=True, arbitrary_types_allowed=True)

    seq: int
    speaker: str
    start_ms: int
    text: str


class ReviewCall(BaseModel):
    """Оценённый звонок в том виде, в каком он уходит в файл."""

    model_config = ConfigDict(frozen=True, arbitrary_types_allowed=True)

    call_id: uuid.UUID
    started_at: datetime
    duration_s: int
    #: Имя сотрудника, а если его нет — логин. Пусто — звонок без сотрудника.
    operator: str
    #: Вердикты по этапам, как их хранит оценка.
    stages: Sequence[Mapping[str, Any]]
    lines: Sequence[ReviewLine]


async def load_lines(
    session: AsyncSession, call_ids: Sequence[uuid.UUID]
) -> dict[uuid.UUID, list[ReviewLine]]:
    """Реплики звонков по порядку. Звонка без реплик в ответе нет."""
    found: dict[uuid.UUID, list[ReviewLine]] = {}
    if not call_ids:
        return found
    rows = await session.execute(_LINES_SQL, {"call_ids": list(call_ids)})
    for row in rows.mappings():
        found.setdefault(row["call_id"], []).append(
            ReviewLine(
                seq=int(row["seq"]),
                speaker=str(row["speaker"]),
                start_ms=int(row["start_ms"]),
                text=str(row["text_raw"]),
            )
        )
    return found


def build_workbook(
    calls: Sequence[ReviewCall], *, script: CallScript, timezone: str, base_url: str
) -> bytes:
    """Собрать файл: «Сверка», «Разговоры» и «Этапы».

    Звонки идут по времени, а не в порядке выборки: номер звонка на «Сверке» и
    в «Разговорах» один, и по нему листы читаются вместе.

    Время — местное время колл-центра без пояса, как в отчёте по звонкам:
    Excel поясов не знает.
    """
    zone = ZoneInfo(timezone)
    ordered = sorted(calls, key=lambda call: (call.started_at, str(call.call_id)))
    workbook = Workbook()
    # Книга создаётся с пустым листом — он не нужен.
    del workbook[workbook.sheetnames[0]]

    check = workbook.create_sheet("Сверка")
    _header(check, CHECK_COLUMNS, highlighted=_CONTROLLER_COLUMNS)
    talks = workbook.create_sheet("Разговоры")
    _header(talks, TALK_COLUMNS)

    # Строки считаются здесь: `max_row` у листа пересчитывается по всем его
    # ячейкам, и на тысячах реплик сборка шла бы квадратично.
    check_row = talk_row = 1
    for number, call in enumerate(ordered, start=1):
        band = _BAND if number % 2 == 0 else None
        moment = call.started_at.astimezone(zone).replace(tzinfo=None)
        who = call.operator or "Без сотрудника"
        by_seq = {line.seq: line for line in call.lines}
        verdicts = {str(item.get("key")): item for item in call.stages if isinstance(item, Mapping)}
        # Этапы — в порядке скрипта, а не в том, как их вернула модель.
        for position, stage in enumerate(script.stages, start=1):
            verdict = verdicts.get(stage.key)
            if verdict is None:
                continue
            check_row += 1
            row = check_row
            _put(check, row, 1, number, band=band)
            _put(check, row, 2, moment, band=band, number_format=_MOMENT_FORMAT)
            _put(check, row, 3, who, band=band)
            _put(check, row, 4, _clock(call.duration_s * 1000), band=band)
            _put_link(check, row, 5, _link(base_url, call.call_id), band=band)
            _put(check, row, 6, f"{position}. {stage.label}", band=band, wrap=True)
            _put(check, row, 7, "да" if stage.mandatory else "", band=band)
            _put(check, row, 8, _mark(verdict), band=band)
            _put(check, row, 9, str(verdict.get("comment") or ""), band=band, wrap=True)
            _put(check, row, 10, _quotes(verdict.get("evidence"), by_seq), band=band, wrap=True)
            _put(check, row, 11, None, band=band)
            _put(check, row, 12, None, band=band, wrap=True)

        for line in call.lines:
            talk_row += 1
            row = talk_row
            _put(talks, row, 1, number, band=band)
            _put(talks, row, 2, moment, band=band, number_format=_MOMENT_FORMAT)
            _put(talks, row, 3, who, band=band)
            _put(talks, row, 4, _clock(line.start_ms), band=band)
            _put(talks, row, 5, _speaker(line.speaker), band=band)
            _put(talks, row, 6, line.text, band=band, wrap=True)

    for sheet in (check, talks):
        sheet.auto_filter.ref = sheet.dimensions

    stages = workbook.create_sheet("Этапы")
    _header(stages, STAGE_COLUMNS)
    for row, (position, stage) in enumerate(enumerate(script.stages, start=1), start=2):
        _put(stages, row, 1, f"{position}. {stage.label}", wrap=True)
        _put(stages, row, 2, "да" if stage.mandatory else "")
        _put(stages, row, 3, stage.requirement, wrap=True)
        _put(stages, row, 4, stage.applies_when or "В любом звонке по скрипту.", wrap=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _mark(verdict: Mapping[str, Any]) -> int | str:
    """Оценка модели так, как её ставит контролёр: 1, 0 или «не понадобился»."""
    if not verdict.get("applicable"):
        return NOT_NEEDED
    return 1 if verdict.get("passed") else 0


def _quotes(evidence: Any, by_seq: Mapping[int, ReviewLine]) -> str:
    """Реплики-основания: время, кто говорит, текст. Номера мимо разговора пропускаются."""
    if not isinstance(evidence, list):
        return ""
    lines = [by_seq[seq] for seq in evidence if isinstance(seq, int) and seq in by_seq]
    return "\n".join(
        f"{_clock(line.start_ms)} {_speaker(line.speaker)}: {line.text}" for line in lines
    )


def _speaker(speaker: str) -> str:
    """Преобразовать роль реплики в подпись участника разговора."""
    return SPEAKER_LABEL.get(speaker, SPEAKER_LABEL["unknown"])


def _clock(milliseconds: int) -> str:
    """Преобразовать время реплики в читаемую подпись для разговора."""
    minutes, seconds = divmod(max(milliseconds, 0) // 1000, 60)
    return f"{minutes:02d}:{seconds:02d}"


def _link(base_url: str, call_id: uuid.UUID) -> str:
    """Окно разбора звонка на экране оценки: там запись и разговор по ролям."""
    base = base_url.strip().rstrip("/")
    return f"{base}/quality?call={call_id}"


def _header(
    sheet: Any, columns: Sequence[tuple[str, int]], *, highlighted: frozenset[int] = frozenset()
) -> None:
    """Добавить и оформить строку заголовков листа отчёта."""
    for index, (name, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        # Через `_put`: в шапке бывают названия из классификатора (выгрузка
        # дашборда), и название, начинающееся с `=`, формулой стать не должно.
        cell = _put(sheet, 1, index, name)
        cell.font = Font(bold=True)
        cell.alignment = _TOP_WRAPPED
        if index in highlighted:
            cell.fill = _TO_FILL
    sheet.freeze_panes = "A2"


def _put(
    sheet: Any,
    row: int,
    column: int,
    value: Any,
    *,
    band: PatternFill | None = None,
    wrap: bool = False,
    number_format: str | None = None,
) -> Any:
    """Записать ячейку. Строка остаётся строкой, даже если похожа на формулу.

    openpyxl записывает строку, начинающуюся с `=`, формулой. Реплики и
    пояснение модели — текст из разговора: формула из них выполнилась бы у
    того, кто откроет файл. Каждая строка проходит `fit_text` — управляющий
    символ иначе ронял бы сборку всего файла.
    """
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = fit_text(value)
        cell.data_type = "s"
    else:
        cell.value = value
    cell.alignment = _TOP_WRAPPED if wrap else _TOP
    if band is not None:
        cell.fill = band
    if number_format:
        cell.number_format = number_format
    return cell


def _put_link(
    sheet: Any, row: int, column: int, url: str, *, band: PatternFill | None = None
) -> None:
    """Ссылка на разбор звонка. Без адреса сервиса — путь текстом: открыть его нечем."""
    if not url.startswith(("http://", "https://")):
        _put(sheet, row, column, url, band=band)
        return
    cell = _put(sheet, row, column, "Открыть", band=band)
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")


__all__ = [
    "NOT_NEEDED",
    "REVIEW_MAX",
    "REVIEW_SIZE",
    "ReviewCall",
    "ReviewLine",
    "build_workbook",
    "load_lines",
]
