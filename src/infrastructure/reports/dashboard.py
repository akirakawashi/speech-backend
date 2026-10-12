"""Числа дашборда файлом Excel.

Три листа. «По времени» — столбцы графика: корзина, звонки по темам экрана,
всего и доля негатива. «Темы» — все темы периода по отдельности: на экране
редкие сложены в «Прочее», в файле место есть. «Подтемы» — уточнения каждой
темы.

Названия тем и подтем — из классификатора, а его правят люди: каждая строка
идёт через `_put`, и название, начинающееся с `=`, формулой не станет.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook

from src.infrastructure.persistence.dashboard import Dashboard
from src.infrastructure.reports.quality import _header, _put

_GRAIN_LABEL = {"day": "День", "week": "Неделя", "month": "Месяц"}
_DAY_FORMAT = "DD.MM.YYYY"
_PERCENT = "0.0%"


def _percent(sheet: Any, row: int, column: int, share: float | None) -> None:
    """Доля — числом с форматом процента: в Excel её можно сортировать и складывать."""
    _put(sheet, row, column, share, number_format=_PERCENT)


def _day(sheet: Any, row: int, column: int, day: date) -> None:
    """Преобразовать календарную дату в ячейку отчёта."""
    _put(sheet, row, column, day, number_format=_DAY_FORMAT)


def build_workbook(board: Dashboard) -> bytes:
    """Построить книгу Excel из готовых данных дашборда."""
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]

    timeline = workbook.create_sheet("По времени")
    _header(
        timeline,
        [
            (f"{_GRAIN_LABEL[board.grain]}, с", 14),
            ("по", 14),
            *((theme.label, 22) for theme in board.themes),
            ("Всего", 12),
            ("Доля негатива", 15),
        ],
    )
    for row, bucket in enumerate(board.buckets, start=2):
        _day(timeline, row, 1, bucket.start)
        _day(timeline, row, 2, bucket.end)
        for offset, calls in enumerate(bucket.by_theme):
            _put(timeline, row, 3 + offset, calls)
        _put(timeline, row, 3 + len(board.themes), bucket.calls)
        _percent(timeline, row, 4 + len(board.themes), bucket.negative_share)

    themes = workbook.create_sheet("Темы")
    _header(themes, [("Тема", 44), ("Обращений", 12), ("Доля", 10), ("Доля негатива", 15)])
    for row, theme in enumerate(board.all_themes, start=2):
        _put(themes, row, 1, theme.label)
        _put(themes, row, 2, theme.calls)
        _percent(themes, row, 3, theme.share)
        _percent(themes, row, 4, theme.negative_share)

    details = workbook.create_sheet("Подтемы")
    _header(details, [("Тема", 44), ("Подтема", 50), ("Обращений", 12), ("Доля в теме", 13)])
    row = 1
    for theme in board.all_themes:
        for subtopic in board.subtopics_by_theme.get(theme.key, ()):
            row += 1
            _put(details, row, 1, theme.label)
            _put(details, row, 2, subtopic.label)
            _put(details, row, 3, subtopic.calls)
            _percent(details, row, 4, subtopic.share)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
