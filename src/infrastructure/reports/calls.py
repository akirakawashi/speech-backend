"""Отчёт по звонкам за период — файлом Excel (SP-604, пять уровней с SP-310).

Для администратора: звонки периода со всеми полями карточки и расшифровкой,
и сводки по операторам, проблемам, причинам обращения,
классификации ИИ, темам по предложению ИИ помощника и совпадению ИИ с оператором.

**Проблема** — классификация, которую оператор выбрал в NCC (`uroven_1…5`),
решение владельца 2026-09-15. Приём пишет её в `speech_calls.disposition` путём
через « › »; звонки, принятые раньше, дозаполняет миграция `sp_0008`.

**Классификация оператора и классификация ИИ — узлы одного дерева** (SP-310,
S-22). Уровни ложатся в колонки по местам, и сравнение идёт уровень в уровень:
процент совпадения — от пяти уровней, пустой уровень оператора или прочерк не
совпадает (решение владельца 2026-09-15, `classifier.level_match`).

**Причина обращения** (S-32) — колонки «Причина обращения» и «Причина по
справочнику» на местах образца владельца (AB, AC) и лист «По причинам
обращения»: оценка модели по расшифровке, а не запись оператора.
Звонок без причины — «[не определена…]», а не код «не удалось понять»: K90
значит «прочли, но не поняли», и смешивать с ним неразобранные нельзя.

**Расшифровка в отчёте есть всегда** (S-31, 2026-09-24; до того SP-604
выгружал отчёт без реплик). Текст — тот же, что в карточке: расшифровка как
есть, без купюр. Нет текста — в ячейке написано почему; пустой ячейки нет.

Отчёт строится **своим запросом по периоду**, а не страницами поиска: у поиска
потолок страницы и выдача под экран, а отчёт — это все звонки периода. Потолок
у отчёта свой — `REPORT_LIMIT`.
"""

from __future__ import annotations

import io
import uuid
from collections import Counter
from collections.abc import Callable, Sequence
from datetime import date, datetime
from typing import Any, Final
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from src.application.dto.calls import CallSource
from src.application.dto.reasons import CallReason
from src.application.period import period_bounds
from src.infrastructure.persistence.classifier import (
    CLASSIFIER_DEPTH,
    NodePaths,
    branch_sql,
    level_match,
    path_caption,
    present_levels,
    suggestion_labels,
)

#: Сколько звонков помещается в один отчёт — около месяца. Замер на узле
#: 2026-09-24: неделя 18–24.09 — 12 330 звонков, выборка 0,9 с, сборка 8,7 с,
#: файл 15,2 МБ. Пятьдесят тысяч — ~40 с сборки и ~60 МБ: меньше минуты, которую
#: nginx по умолчанию ждёт молчащий сервер (настроек прокси перед узлом мы не
#: знаем). Сто тысяч, как было до расшифровки, — ~80 с, уже нет. Больше — сузить
#: период.
REPORT_LIMIT: Final[int] = 50_000

#: Предел ячейки Excel в знаках. Длиннее openpyxl режет молча — поэтому режем
#: сами и с пометкой, которая помещается внутри предела.
EXCEL_CELL_LIMIT: Final[int] = 32_767
TRUNCATED_MARK: Final[str] = """
[… обрезано: превышен лимит ячейки Excel, полный текст — в карточке звонка]"""

#: Подпись реплики в расшифровке. `unknown` — честно «не определено»: у
#: моно-записи роли может не быть, и назвать реплику клиентской было бы выдумкой.
SPEAKER_LABEL: Final[dict[str, str]] = {
    "operator": "Оператор",
    "client": "Клиент",
    "unknown": "Не определено",
}

#: Почему у звонка нет текста — по статусу обработки. Остальные статусы значат,
#: что звонок ещё идёт по конвейеру.
TRANSCRIPT_NOTE: Final[dict[str, str]] = {
    "skipped_short": "[звонок короче порога расшифровки]",
    "skipped_no_audio": "[записи нет у источника]",
    "failed": "[ошибка обработки — расшифровки нет]",
    "done": "[расшифровки нет]",
}
TRANSCRIPT_NO_SPEECH: Final[str] = "[в записи нет речи]"
TRANSCRIPT_IN_PROGRESS: Final[str] = "[звонок в обработке — расшифровки ещё нет]"

#: Причины нет. Два разных состояния, и оба — не код справочника.
REASON_NOT_ANALYZED: Final[str] = "[не определена: звонок не разобран]"
REASON_UNDETERMINED: Final[str] = "[не определена]"

SENTIMENT_LABEL: Final[dict[str, str]] = {
    "negative": "негативный",
    "neutral": "нейтральный",
    "positive": "позитивный",
}
SEGMENT_LABEL: Final[dict[str, str]] = {"b2b": "юрлицо", "individual": "физлицо"}
PRODUCT_LABEL: Final[dict[str, str]] = {
    "wallet": "Кошелёк",
    "mobile": "Мобильная связь",
    "broadband": "ШПД",
}
SOURCE_LABEL: Final[dict[str, str]] = {"project": "из очереди", "llm": "определено моделью"}

_LEVELS: Final[range] = range(1, CLASSIFIER_DEPTH + 1)

#: Колонки листа «Звонки» и их ширина. Уровни — по местам: пятая колонка
#: оператора и пятая колонка ИИ — один и тот же уровень дерева.
CALL_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("Начало", 20),
    ("Окончание", 20),
    ("Длительность, с", 15),
    ("Оператор", 18),
    ("ФИО оператора", 24),
    ("Очередь", 16),
    ("Проблема (классификация NCC)", 55),
    *((f"Уровень {level}: оператор", 24) for level in _LEVELS),
    *((f"Уровень {level}: ИИ", 24) for level in _LEVELS),
    ("Совпало уровней", 16),
    ("Совпадение, %", 15),
    ("Тема по предложению ИИ помощника", 55),
    ("Новая формулировка", 12),
    ("Тональность", 14),
    ("Сегмент", 10),
    ("Сегмент: откуда", 20),
    ("Продукт", 16),
    ("Продукт: откуда", 20),
    ("Краткое содержание", 80),
    #: AB и AC — как в образце владельца от 2026-09-24 (S-32).
    ("Причина обращения", 55),
    ("Причина по справочнику", 45),
    ("Обработка", 12),
    ("Анализ", 10),
    ("Запись", 12),
    ("Идентификатор сессии", 44),
    ("Ссылка", 60),
    #: Последней: причины обращения стоят на местах образца владельца (AB, AC),
    #: и его формулы `COUNTIF(Звонки!$AC:$AC, …)` не съезжают.
    ("Расшифровка", 100),
)

#: Колонки сводных листов после первой.
SUMMARY_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("Звонков", 10),
    ("Доля звонков, %", 16),
    ("Средняя длительность, с", 24),
    ("Негативных", 12),
    ("Нейтральных", 13),
    ("Позитивных", 12),
    ("Совпадение ИИ с оператором, %", 30),
)

#: Колонки листа «Совпадение с оператором».
MATCH_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("Срез", 40),
    ("Сравнено", 12),
    ("Совпало", 12),
    ("Совпадение, %", 15),
    ("Что считается", 70),
)

#: Условия фильтров — одни у подсчёта и у выборки: разойдясь, они дали бы отказ
#: «слишком много» на отчёт, который поместился бы.
#:
#: Узел классификатора отбирает всю свою ветку: выбрав «Физ.лицо › ШПД», человек
#: спрашивает про все звонки по ШПД, а не про те, где классификация не пошла
#: глубже.
_FILTERS_SQL = f"""
       c.started_at >= :date_from
   AND c.started_at <  :date_to
   AND c.source = :source
   AND (CAST(:operator_id AS uuid) IS NULL OR c.operator_id = CAST(:operator_id AS uuid))
   AND {branch_sql("c.operator_node_id", "operator_node_id")}
   AND {branch_sql("c.ai_node_id", "ai_node_id")}
"""

_COUNT_SQL = text(f"SELECT count(*) FROM speech_calls AS c WHERE {_FILTERS_SQL}")  # noqa: S608

_SPEAKER_SQL = " ".join(
    f"WHEN '{speaker}' THEN '{label}: '" for speaker, label in SPEAKER_LABEL.items()
)

#: Расшифровка собирается из реплик здесь же: склеенного текста у звонка нет
#: намеренно (`SpeechTranscript`). Расшифровка без реплик — пустая строка, не
#: NULL: «речи нет» отличается от «ещё не готово».
_ROWS_SQL = text(
    f"""
    SELECT c.id, c.session_id, c.started_at, c.ended_at, c.duration_s,
           o.login AS operator_login, o.full_name AS operator_name,
           p.title AS project_title, c.disposition,
           c.operator_node_id, c.ai_node_id, c.ai_suggestion, c.ai_suggestion_node_id,
           c.sentiment, c.segment, c.segment_source, c.product, c.product_source,
           c.summary, c.processing_status, c.analysis_status, c.audio_status,
           c.call_reason, c.call_reason_code, r.label AS call_reason_label,
           CASE WHEN t.id IS NOT NULL THEN COALESCE(
               (SELECT string_agg(
                           CASE s.speaker {_SPEAKER_SQL} ELSE 'Не определено: ' END
                           || s.text_raw,
                           E'\\n' ORDER BY s.seq)
                  FROM speech_transcript_segments AS s
                 WHERE s.call_id = c.id),
               '') END AS transcript
      FROM speech_calls AS c
      LEFT JOIN speech_operators AS o ON o.id = c.operator_id
      LEFT JOIN speech_ncc_projects AS p ON p.id = c.project_id
      LEFT JOIN speech_transcripts AS t ON t.call_id = c.id
      LEFT JOIN speech_call_reasons AS r ON r.code = c.call_reason_code
     WHERE {_FILTERS_SQL}
     ORDER BY c.started_at, c.id
    """  # noqa: S608
)

#: Справочник причин целиком: в сводке — все включённые коды, в том числе ни разу
#: не встретившиеся, и выключенные, если звонки периода ими размечены.
_REASONS_SQL = text(
    "SELECT code, label, is_active FROM speech_call_reasons ORDER BY sort_order, code"
)

#: Звонки периода по узлам — классификации оператора и классификации ИИ. В ветки
#: их сворачивает `NodePaths.rollup`.
_CLASSIFICATION_SQL = text(
    """
    SELECT 'operator' AS side, c.operator_node_id AS node_id, count(*) AS calls
      FROM speech_calls AS c
     WHERE c.started_at >= :date_from
       AND c.started_at <  :date_to
       AND c.source = :source
       AND c.operator_node_id IS NOT NULL
     GROUP BY c.operator_node_id
    UNION ALL
    SELECT 'ai' AS side, c.ai_node_id AS node_id, count(*) AS calls
      FROM speech_calls AS c
     WHERE c.started_at >= :date_from
       AND c.started_at <  :date_to
       AND c.source = :source
       AND c.ai_node_id IS NOT NULL
     GROUP BY c.ai_node_id
    """
)


class ReportTooLarge(Exception):
    """В период попало больше звонков, чем помещается в один отчёт."""

    def __init__(self, rows: int, limit: int) -> None:
        """Настроить ReportTooLarge из переданных параметров и зависимостей."""
        super().__init__(f"{rows} calls in the period, the report takes {limit}")
        self.rows = rows
        self.limit = limit


class ReportFilters(BaseModel):
    """Проверенные параметры выборки звонков для отчёта."""

    model_config = ConfigDict(frozen=True, arbitrary_types_allowed=True)

    date_from: datetime
    date_to: datetime
    operator_id: uuid.UUID | None = None
    #: Ветки классификатора — по классификации оператора и по классификации ИИ.
    operator_node_id: uuid.UUID | None = None
    ai_node_id: uuid.UUID | None = None
    #: Источник звонков (S-35): отчёт NCC и отчёт Телфина — разные отчёты, и
    #: звонки одного в другой не попадают.
    source: str = CallSource.NCC.value

    def params(self) -> dict[str, Any]:
        """Вернуть именованные параметры SQL из проверенных фильтров отчёта."""
        return {
            "date_from": self.date_from,
            "date_to": self.date_to,
            "source": self.source,
            "operator_id": self.operator_id,
            "operator_node_id": self.operator_node_id,
            "ai_node_id": self.ai_node_id,
        }


class ReportRow(BaseModel):
    """Строка листа «Звонки»: всё из карточки звонка и расшифровка."""

    model_config = ConfigDict(frozen=True, arbitrary_types_allowed=True)

    call_id: uuid.UUID
    session_id: str
    started_at: datetime
    ended_at: datetime
    duration_s: int
    operator_login: str | None
    #: ФИО, если справочник его знает. Логин вместо имени сюда не пишется: у
    #: оператора, которого завёл приём, имени нет, и повторять логин незачем.
    operator_name: str | None
    project: str | None
    problem: str | None
    sentiment: str | None
    segment: str | None
    segment_source: str
    product: str | None
    product_source: str
    summary: str | None
    processing_status: str
    analysis_status: str
    audio_status: str
    #: Пути от корня: классификация оператора и классификация ИИ (SP-310).
    operator_path: tuple[str, ...] = ()
    ai_path: tuple[str, ...] = ()
    #: Тема по предложению ИИ помощника — путь; `suggestion_is_new` — со своей
    #: формулировкой, которой в дереве нет.
    suggestion: tuple[str, ...] = ()
    suggestion_is_new: bool = False
    #: Расшифровка «Оператор: … / Клиент: …» по порядку реплик.
    #: `None` — расшифровки нет; `""` — есть, но речи в записи нет.
    transcript: str | None = None
    #: Причина обращения (S-32): фраза, код и подпись кода из справочника.
    call_reason: str | None = None
    call_reason_code: str | None = None
    call_reason_label: str | None = None

    @property
    def reason_caption(self) -> str:
        """«K05 · Не работает домашний интернет» — или почему причины нет."""
        if self.call_reason_code:
            return CallReason(
                code=self.call_reason_code, label=self.call_reason_label or ""
            ).caption
        if self.analysis_status != "done":
            return REASON_NOT_ANALYZED
        return REASON_UNDETERMINED

    @property
    def transcript_cell(self) -> str:
        """Текст для ячейки: расшифровка или пометка, почему её нет — пустой ячейки нет."""
        if self.transcript:
            return self.transcript
        if self.transcript == "":
            return TRANSCRIPT_NO_SPEECH
        return TRANSCRIPT_NOTE.get(self.processing_status, TRANSCRIPT_IN_PROGRESS)

    @property
    def matches(self) -> tuple[bool, ...]:
        """Совпадение ИИ с оператором по пяти уровням; пусто — сравнивать не с чем."""
        return level_match(self.operator_path, self.ai_path)


class ClassificationOption(BaseModel):
    """Узел, который в периоде встречается, — вариант фильтра отчёта."""

    model_config = ConfigDict(frozen=True, arbitrary_types_allowed=True)

    node_id: uuid.UUID
    parent_id: uuid.UUID | None
    level: int
    label: str
    sort_order: int
    #: Звонков в ветке узла — по классификации оператора и по классификации ИИ.
    operator_calls: int
    ai_calls: int


async def load_report(
    session: AsyncSession, filters: ReportFilters, *, limit: int
) -> list[ReportRow]:
    """Звонки отчёта. Больше `limit` — отказ до выборки, а не обрезанный файл.

    Обрезанный отчёт выглядел бы полным: сводка на его листах посчитала бы
    часть периода и назвала бы это периодом.
    """
    params = filters.params()
    total = int(await session.scalar(_COUNT_SQL, params) or 0)
    if total > limit:
        raise ReportTooLarge(total, limit)
    rows = (await session.execute(_ROWS_SQL, params)).mappings().all()
    paths = await NodePaths.load(session) if rows else NodePaths(())
    return [_report_row(row, paths) for row in rows]


async def list_classification(
    session: AsyncSession,
    date_from: datetime,
    date_to: datetime,
    source: str = CallSource.NCC.value,
) -> list[ClassificationOption]:
    """Узлы, которые в периоде встречаются, со счётом звонков в их ветках — в порядке дерева.

    Счёт по ветке, а не по узлу: фильтр отбирает ветку, и число рядом с узлом
    обязано сходиться с тем, что получит выбравший его.
    """
    found: dict[str, dict[uuid.UUID, int]] = {"operator": {}, "ai": {}}
    params = {"date_from": date_from, "date_to": date_to, "source": source}
    for row in (await session.execute(_CLASSIFICATION_SQL, params)).mappings():
        found[str(row["side"])][row["node_id"]] = int(row["calls"])
    if not found["operator"] and not found["ai"]:
        return []
    paths = await NodePaths.load(session)
    operator, ai = paths.rollup(found["operator"]), paths.rollup(found["ai"])
    return [
        ClassificationOption(
            node_id=node,
            parent_id=paths.node(node)["parent_id"],
            level=int(paths.node(node)["level"]),
            label=str(paths.node(node)["label"]),
            sort_order=int(paths.node(node)["sort_order"] or 0),
            operator_calls=operator.get(node, 0),
            ai_calls=ai.get(node, 0),
        )
        for node in sorted(set(operator) | set(ai), key=paths.order)
    ]


def _report_row(row: Any, paths: NodePaths) -> ReportRow:
    """Собрать строку отчёта из данных звонка и снимка дерева."""
    name = row["operator_name"]
    suggestion = suggestion_labels(row["ai_suggestion"])
    return ReportRow(
        call_id=row["id"],
        session_id=str(row["session_id"]),
        started_at=row["started_at"],
        ended_at=row["ended_at"],
        duration_s=int(row["duration_s"] or 0),
        operator_login=row["operator_login"],
        operator_name=name if name and name != row["operator_login"] else None,
        project=row["project_title"],
        problem=row["disposition"],
        sentiment=row["sentiment"],
        segment=row["segment"],
        segment_source=str(row["segment_source"]),
        product=row["product"],
        product_source=str(row["product_source"]),
        summary=row["summary"],
        processing_status=str(row["processing_status"]),
        analysis_status=str(row["analysis_status"]),
        audio_status=str(row["audio_status"]),
        operator_path=paths.labels(row["operator_node_id"]),
        ai_path=paths.labels(row["ai_node_id"]),
        suggestion=suggestion,
        suggestion_is_new=bool(suggestion) and row["ai_suggestion_node_id"] is None,
        transcript=row["transcript"],
        call_reason=row["call_reason"],
        call_reason_code=row["call_reason_code"],
        call_reason_label=row["call_reason_label"],
    )


async def load_reasons(session: AsyncSession) -> list[tuple[CallReason, bool]]:
    """Справочник причин с признаком «включён» — для листа «По причинам обращения»."""
    return [
        (CallReason(code=str(row["code"]), label=str(row["label"])), bool(row["is_active"]))
        for row in (await session.execute(_REASONS_SQL)).mappings()
    ]


def transcript_counts(rows: Sequence[ReportRow]) -> tuple[int, int]:
    """Сколько звонков ушло в файл с текстом и сколько из них обрезано — для журнала."""
    texts = [fit_text(row.transcript) for row in rows if row.transcript]
    return len(texts), sum(text.endswith(TRUNCATED_MARK) for text in texts)


def fit_text(value: str) -> str:
    """Строка, которую Excel примет целиком: без управляющих символов и в пределе ячейки.

    Управляющий символ openpyxl не пропускает — `IllegalCharacterError` с
    **всем значением** в тексте исключения: отчёт отдал бы 500, а расшифровка
    легла бы в журнал ошибок. Длиннее предела — обрезка с пометкой внутри него.
    """
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) <= EXCEL_CELL_LIMIT:
        return value
    return value[: EXCEL_CELL_LIMIT - len(TRUNCATED_MARK)] + TRUNCATED_MARK


class Group(BaseModel):
    """Одна строка сводки."""

    model_config = ConfigDict(frozen=False, arbitrary_types_allowed=True)

    calls: int = 0
    duration_s: int = 0
    sentiments: Counter[str] = Field(default_factory=Counter)
    #: Звонков, где было что сравнить, и сколько уровней у них совпало.
    compared: int = 0
    matched_levels: int = 0

    @property
    def average_duration_s(self) -> int:
        """Вернуть среднюю длительность выбранных разговоров в секундах."""
        return round(self.duration_s / self.calls) if self.calls else 0

    @property
    def match_percent(self) -> float | None:
        """Совпавших уровней от пяти на каждый сравненный звонок; сравнивать не с чем — пусто."""
        return _percent(self.matched_levels, self.compared * CLASSIFIER_DEPTH)


def summarize(
    rows: Sequence[ReportRow], key: Callable[[ReportRow], str | None], empty: str
) -> list[tuple[str, Group]]:
    """Сводка по признаку: больше звонков — выше; при равном счёте — по названию,
    и звонки без значения — последними среди равных.

    Звонки без значения не выбрасываются, а собираются под `empty`: сводка, в
    которой строк меньше, чем звонков в отчёте, читается как ошибка счёта.
    """
    groups: dict[str, Group] = {}
    for row in rows:
        group = groups.setdefault(key(row) or empty, Group())
        group.calls += 1
        group.duration_s += row.duration_s
        if row.sentiment:
            group.sentiments[row.sentiment] += 1
        marks = row.matches
        if marks:
            group.compared += 1
            group.matched_levels += sum(marks)
    return sorted(groups.items(), key=lambda item: (-item[1].calls, item[0] == empty, item[0]))


def build_workbook(
    rows: Sequence[ReportRow],
    *,
    period: tuple[date, date],
    filters: Sequence[tuple[str, str]],
    timezone: str,
    base_url: str,
    generated_at: datetime,
    reasons: Sequence[tuple[CallReason, bool]] = (),
) -> bytes:
    """Собрать файл: «Отчёт», «Звонки», пять сводок и совпадение с оператором.

    `reasons` — справочник причин с признаком «включён» (`load_reasons`): сводка
    по причинам перечисляет все включённые коды, и нулевые тоже, — иначе «ни
    одного звонка с этой причиной» не отличить от «такой причины нет».

    Время — местное время колл-центра без пояса: Excel поясов не знает, и
    `2026-09-13 22:24 UTC` в ячейке читалось бы как вечер 13-го, хотя звонок был
    ночью 14-го.
    """
    zone = ZoneInfo(timezone)
    workbook = Workbook(write_only=True)

    about = workbook.create_sheet("Отчёт")
    about.column_dimensions["A"].width = 30
    about.column_dimensions["B"].width = 60
    for name, value in (
        ("Отчёт", "Звонки речевой аналитики"),
        ("Период с", period[0]),
        ("Период по", period[1]),
        *filters,
        ("Звонков", len(rows)),
        ("Сформирован", _local(generated_at, zone)),
        ("Часовой пояс", timezone),
        (
            "Расшифровка",
            "текст расшифровки как есть, без купюр. Длиннее 32 767 знаков — обрезана с"
            " пометкой; полный текст — в карточке звонка",
        ),
        (
            "Совпадение, %",
            "совпавших уровней из пяти; пустой уровень оператора или прочерк не совпадает",
        ),
        (
            "Причина обращения",
            "оценка модели по расшифровке — зачем звонил клиент, а не запись"
            " оператора; код — из справочника причин контакт-центра",
        ),
    ):
        about.append([_bold(about, name), _cell(about, value)])

    calls = workbook.create_sheet("Звонки")
    _header(calls, CALL_COLUMNS)
    for row in rows:
        marks = row.matches
        matched = sum(marks) if marks else None
        calls.append(
            [
                _cell(calls, value)
                for value in (
                    _local(row.started_at, zone),
                    _local(row.ended_at, zone),
                    row.duration_s,
                    row.operator_login,
                    row.operator_name,
                    row.project,
                    row.problem,
                    *_levels(row.operator_path),
                    *_levels(row.ai_path),
                    matched,
                    _percent(matched, CLASSIFIER_DEPTH) if matched is not None else None,
                    path_caption(row.suggestion) or None,
                    ("да" if row.suggestion_is_new else "нет") if row.suggestion else None,
                    SENTIMENT_LABEL.get(row.sentiment or "", row.sentiment),
                    SEGMENT_LABEL.get(row.segment or "", row.segment),
                    SOURCE_LABEL.get(row.segment_source) if row.segment else None,
                    PRODUCT_LABEL.get(row.product or "", row.product),
                    SOURCE_LABEL.get(row.product_source) if row.product else None,
                    row.summary,
                    row.call_reason,
                    row.reason_caption,
                    row.processing_status,
                    row.analysis_status,
                    row.audio_status,
                    row.session_id,
                    _link(base_url, row.call_id),
                    row.transcript_cell,
                )
            ]
        )
    calls.auto_filter.ref = f"A1:{get_column_letter(len(CALL_COLUMNS))}{len(rows) + 1}"

    new_wordings = {path_caption(row.suggestion) for row in rows if row.suggestion_is_new}
    for title, first, key, empty in (
        ("По операторам", "Оператор", lambda r: r.operator_login, "не указан"),
        ("По проблемам", "Проблема (классификация NCC)", lambda r: r.problem, "не указана"),
        ("По классификации ИИ", "Классификация ИИ", lambda r: path_caption(r.ai_path), "нет"),
        (
            "По темам ИИ помощника",
            "Тема по предложению ИИ помощника",
            lambda r: path_caption(r.suggestion),
            "нет",
        ),
    ):
        if title == "По классификации ИИ":
            # Пятым, как в образце владельца: сразу за проблемами оператора.
            _reasons_sheet(workbook, rows, reasons)
        sheet = workbook.create_sheet(title)
        is_ai = title == "По темам ИИ помощника"
        columns = ((first, 55), *SUMMARY_COLUMNS, *((("Новая формулировка", 20),) if is_ai else ()))
        _header(sheet, columns)
        for name, group in summarize(rows, key, empty):
            values = _summary_values(name, group, len(rows))
            if is_ai:
                values.append("да" if name in new_wordings else "")
            sheet.append([_cell(sheet, value) for value in values])

    _match_sheet(workbook, rows)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _summary_values(name: str, group: Group, total: int) -> list[Any]:
    """Рассчитать значения итоговой строки отчёта."""
    return [
        name,
        group.calls,
        round(group.calls / total * 100, 1) if total else 0,
        group.average_duration_s,
        group.sentiments["negative"],
        group.sentiments["neutral"],
        group.sentiments["positive"],
        group.match_percent,
    ]


def _reasons_sheet(
    workbook: Workbook,
    rows: Sequence[ReportRow],
    reasons: Sequence[tuple[CallReason, bool]],
) -> None:
    """Звонки периода по причинам обращения — готовыми числами, не формулами.

    Встретившиеся причины — по убыванию числа звонков; затем включённые коды, у
    которых звонков нет, — в порядке справочника. Звонки без причины —
    отдельными строками: сумма по листу равна числу звонков отчёта.
    """
    sheet = workbook.create_sheet("По причинам обращения")
    _header(sheet, (("Причина обращения (справочник)", 70), *SUMMARY_COLUMNS))
    groups = summarize(rows, lambda r: r.reason_caption, REASON_UNDETERMINED)
    seen = {name for name, _ in groups}
    for name, group in groups:
        sheet.append([_cell(sheet, value) for value in _summary_values(name, group, len(rows))])
    for reason, is_active in reasons:
        if is_active and reason.caption not in seen:
            sheet.append(
                [
                    _cell(sheet, value)
                    for value in _summary_values(reason.caption, Group(), len(rows))
                ]
            )


def _match_sheet(workbook: Workbook, rows: Sequence[ReportRow]) -> None:
    """Совпадение ИИ с оператором: по каждому уровню и по правилу владельца в целом.

    Звонки без сравнения — отдельной строкой: сводка, в которой звонков меньше,
    чем в отчёте, читается как ошибка счёта.
    """
    sheet = workbook.create_sheet("Совпадение с оператором")
    _header(sheet, MATCH_COLUMNS)
    compared = [row for row in rows if row.matches]
    for index in range(CLASSIFIER_DEPTH):
        with_level = [row for row in compared if present_levels(row.operator_path)[index]]
        matched = sum(row.matches[index] for row in with_level)
        sheet.append(
            [
                _cell(sheet, value)
                for value in (
                    f"Уровень {index + 1}",
                    len(with_level),
                    matched,
                    _percent(matched, len(with_level)),
                    f"звонки, где у оператора есть уровень {index + 1}; совпало — звонков",
                )
            ]
        )
    levels = sum(sum(row.matches) for row in compared)
    for values in (
        (
            "Все пять уровней",
            len(compared) * CLASSIFIER_DEPTH,
            levels,
            _percent(levels, len(compared) * CLASSIFIER_DEPTH),
            "уровни сравненных звонков, по пять на звонок, — правило владельца",
        ),
        (
            "Звонков сравнено",
            len(compared),
            None,
            None,
            "у звонка есть и классификация оператора, и классификация ИИ",
        ),
        (
            "Звонков без сравнения",
            len(rows) - len(compared),
            None,
            None,
            "нет классификации оператора или классификации ИИ",
        ),
    ):
        sheet.append([_cell(sheet, value) for value in values])


def _levels(path: Sequence[str]) -> list[str | None]:
    """Уровни пути по местам. Пустой уровень и уровень за концом пути — пустая ячейка."""
    return [
        (path[index] or None) if index < len(path) else None for index in range(CLASSIFIER_DEPTH)
    ]


def _percent(part: int, whole: int) -> float | None:
    """Форматировать долю в процентах; пустой знаменатель обозначить прочерком."""
    return round(part / whole * 100, 1) if whole else None


def _local(moment: datetime | None, zone: ZoneInfo) -> datetime | None:
    """Преобразовать момент звонка в часовой пояс отчёта."""
    return moment.astimezone(zone).replace(tzinfo=None) if moment else None


def _link(base_url: str, call_id: uuid.UUID) -> str:
    """Подготовить ссылку на карточку звонка в отчёте."""
    base = base_url.strip().rstrip("/")
    return f"{base}/calls/{call_id}" if base else f"/calls/{call_id}"


def _cell(sheet: Any, value: Any) -> Any:
    """Значение ячейки. Строка остаётся строкой, даже если похожа на формулу.

    openpyxl записывает строку, начинающуюся с `=`, формулой. Краткое содержание,
    формулировки темы и расшифровка — текст из разговора, то есть их содержимое
    нам не принадлежит, и формула из них выполнилась бы у того, кто откроет файл.
    Каждая строка проходит `fit_text`: один управляющий символ иначе ронял бы
    весь отчёт.
    """
    if not isinstance(value, str):
        return value
    cell = WriteOnlyCell(sheet, value=fit_text(value))
    cell.data_type = "s"
    return cell


def _bold(sheet: Any, value: str) -> Any:
    """Оформить выделенную ячейку листа отчёта."""
    cell = _cell(sheet, value)
    cell.font = Font(bold=True)
    return cell


def _header(sheet: Any, columns: Sequence[tuple[str, int]]) -> None:
    """Добавить и оформить строку заголовков листа отчёта."""
    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.append([_bold(sheet, name) for name, _ in columns])


__all__ = [
    "REPORT_LIMIT",
    "ClassificationOption",
    "Group",
    "ReportFilters",
    "ReportRow",
    "ReportTooLarge",
    "build_workbook",
    "fit_text",
    "list_classification",
    "load_reasons",
    "load_report",
    "period_bounds",
    "summarize",
    "transcript_counts",
]
