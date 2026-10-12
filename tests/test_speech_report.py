"""Отчёт в Excel — сборка файла, без базы (SP-604, пять уровней с SP-310).

Проверяется то, что увидит человек, открыв файл: листы, колонки, местное время,
уровни по местам, совпадение с оператором по правилу владельца, сводки и то, что
текст из разговора не выполнится формулой.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime
from typing import Any

import pytest
from openpyxl import load_workbook

from src.domain.calls.reasons import CallReason
from src.infrastructure.reports.calls import (
    EXCEL_CELL_LIMIT,
    TRUNCATED_MARK,
    ReportRow,
    build_workbook,
    fit_text,
    summarize,
    transcript_counts,
)

GENERATED = datetime(2026, 9, 15, 12, 0, tzinfo=UTC)
OPERATOR = ("Физ.лицо", "Мобильная связь", "Консультация", "Тарифы", "Смена тарифа")
AI = ("Физ.лицо", "Мобильная связь", "Консультация", "Тарифы", "Подключение услуг")


def _row(**overrides: Any) -> ReportRow:
    values: dict[str, Any] = {
        "call_id": uuid.UUID("b663725d-b3d8-4a96-8f0d-07c349ae6f5a"),
        "session_id": "node_domain_0_nauss_0_1778844752_17265555",
        "started_at": datetime(2026, 9, 13, 22, 24, 15, tzinfo=UTC),
        "ended_at": datetime(2026, 9, 13, 22, 33, 7, tzinfo=UTC),
        "duration_s": 532,
        "operator_login": "operator1",
        "operator_name": None,
        "project": "Волна",
        "problem": "Физ.лицо › Мобильная связь › Консультация › Тарифы › Смена тарифа",
        "sentiment": "negative",
        "segment": "individual",
        "segment_source": "llm",
        "product": "mobile",
        "product_source": "llm",
        "summary": "Абонент оспаривал списание 100 рублей.",
        "processing_status": "done",
        "analysis_status": "done",
        "audio_status": "stored",
        "operator_path": OPERATOR,
        "ai_path": AI,
        "suggestion": ("Физ.лицо", "Мобильная связь", "Списание за подписку"),
        "suggestion_is_new": True,
    }
    values.update(overrides)
    return ReportRow(**values)


def _build(
    rows: list[ReportRow],
    base_url: str = "",
    reasons: tuple[tuple[CallReason, bool], ...] = (),
) -> Any:
    content = build_workbook(
        rows,
        period=(date(2026, 9, 14), date(2026, 9, 14)),
        filters=[("Классификация ИИ (с веткой)", "Физ.лицо › Мобильная связь")],
        timezone="Europe/Moscow",
        base_url=base_url,
        generated_at=GENERATED,
        reasons=reasons,
    )
    return load_workbook(io.BytesIO(content))


def _table(sheet: Any) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, row, strict=True)) for row in rows[1:]]


def _levels(call: dict[str, Any], side: str) -> list[Any]:
    return [call[f"Уровень {level}: {side}"] for level in range(1, 6)]


def test_the_file_has_the_calls_the_summaries_and_the_match() -> None:
    workbook = _build([_row()])

    assert workbook.sheetnames == [
        "Отчёт",
        "Звонки",
        "По операторам",
        "По проблемам",
        "По причинам обращения",
        "По классификации ИИ",
        "По темам ИИ помощника",
        "Совпадение с оператором",
    ]


def test_a_call_row_carries_every_field_in_the_call_centre_time() -> None:
    """Время — местное без пояса: Excel поясов не знает, а звонок был ночью 14-го."""
    (call,) = _table(_build([_row()])["Звонки"])

    assert call["Начало"] == datetime(2026, 9, 14, 1, 24, 15)
    assert call["Окончание"] == datetime(2026, 9, 14, 1, 33, 7)
    assert call["Длительность, с"] == 532
    assert call["Оператор"] == "operator1"
    assert call["ФИО оператора"] is None, "логин вместо имени не повторяется"
    assert call["Очередь"] == "Волна"
    assert call["Проблема (классификация NCC)"].startswith("Физ.лицо › Мобильная связь")
    assert (
        call["Тема по предложению ИИ помощника"]
        == "Физ.лицо › Мобильная связь › Списание за подписку"
    )
    assert call["Новая формулировка"] == "да"
    assert (call["Тональность"], call["Сегмент"], call["Сегмент: откуда"]) == (
        "негативный",
        "физлицо",
        "определено моделью",
    )
    assert call["Краткое содержание"] == "Абонент оспаривал списание 100 рублей."
    assert call["Ссылка"] == "/calls/b663725d-b3d8-4a96-8f0d-07c349ae6f5a"


def test_both_paths_lie_level_by_level_and_the_match_counts_five_levels() -> None:
    """Пятая колонка оператора и пятая колонка ИИ — один и тот же уровень дерева."""
    (call,) = _table(_build([_row()])["Звонки"])

    assert _levels(call, "оператор") == list(OPERATOR)
    assert _levels(call, "ИИ") == list(AI)
    assert (call["Совпало уровней"], call["Совпадение, %"]) == (4, 80)


def test_an_empty_level_keeps_its_column_and_never_matches() -> None:
    """«Кошелёк» без уровня 1: уровни не съезжают, а пустой уровень оператора не совпадает."""
    wallet = ("", "Кошелек", "Консультация", "Финансы")
    (call,) = _table(_build([_row(operator_path=wallet, ai_path=wallet)])["Звонки"])

    assert _levels(call, "оператор") == [None, "Кошелек", "Консультация", "Финансы", None]
    assert (call["Совпало уровней"], call["Совпадение, %"]) == (3, 60)


def test_a_call_without_one_of_the_paths_is_not_compared() -> None:
    (call,) = _table(_build([_row(ai_path=(), suggestion=(), suggestion_is_new=False)])["Звонки"])

    assert (call["Совпало уровней"], call["Совпадение, %"]) == (None, None)
    assert (call["Тема по предложению ИИ помощника"], call["Новая формулировка"]) == (None, None)


def test_the_match_sheet_counts_each_level_and_the_whole_rule() -> None:
    rows = [
        _row(),
        _row(
            operator_path=("", "Кошелек", "Консультация", "Финансы"),
            ai_path=("", "Кошелек", "Консультация", "Финансы", "Баланс"),
        ),
        _row(ai_path=()),
    ]

    match = {row["Срез"]: row for row in _table(_build(rows)["Совпадение с оператором"])}

    assert [
        (match[f"Уровень {level}"]["Сравнено"], match[f"Уровень {level}"]["Совпало"])
        for level in range(1, 6)
    ] == [(1, 1), (2, 2), (2, 2), (2, 2), (1, 0)], (
        "пустой уровень 1 и отсутствующий уровень 5 оператора в срез уровня не входят"
    )
    assert (match["Уровень 1"]["Совпадение, %"], match["Уровень 5"]["Совпадение, %"]) == (100, 0)
    whole = match["Все пять уровней"]
    assert (whole["Сравнено"], whole["Совпало"], whole["Совпадение, %"]) == (10, 7, 70)
    assert match["Звонков сравнено"]["Сравнено"] == 2
    assert match["Звонков без сравнения"]["Сравнено"] == 1, "звонок без пути ИИ не потерян"


def test_the_about_sheet_names_the_period_filters_and_what_the_transcript_is() -> None:
    about = {row[0]: row[1] for row in _build([_row()])["Отчёт"].iter_rows(values_only=True)}

    assert about["Период с"] == datetime(2026, 9, 14, 0, 0)
    assert about["Классификация ИИ (с веткой)"] == "Физ.лицо › Мобильная связь"
    assert about["Звонков"] == 1
    assert about["Сформирован"] == datetime(2026, 9, 15, 15, 0)
    assert "как есть, без купюр" in about["Расшифровка"]
    assert "Реплики" not in about, "прежняя строка «не выгружаются» больше не правда"
    assert "из пяти" in about["Совпадение, %"]


def test_the_transcript_is_the_last_column_so_the_reason_columns_keep_their_letters() -> None:
    """Образец владельца считает причины формулой по колонке AC — она не должна съехать."""
    text = "Оператор: Здравствуйте.\nКлиент: Не работает интернет, номер 89780001122."
    sheet = _build([_row(transcript=text)])["Звонки"]
    header = [cell.value for cell in sheet[1]]

    assert header[-1] == "Расшифровка"
    assert sheet.cell(row=2, column=len(header)).value == text


@pytest.mark.parametrize(
    ("overrides", "expected"),
    [
        ({"transcript": ""}, "[в записи нет речи]"),
        ({"processing_status": "skipped_short"}, "[звонок короче порога расшифровки]"),
        ({"processing_status": "skipped_no_audio"}, "[записи нет у источника]"),
        ({"processing_status": "failed"}, "[ошибка обработки — расшифровки нет]"),
        ({"processing_status": "transcribing"}, "[звонок в обработке — расшифровки ещё нет]"),
    ],
)
def test_a_call_without_text_says_why(overrides: dict[str, Any], expected: str) -> None:
    (call,) = _table(_build([_row(**overrides)])["Звонки"])

    assert call["Расшифровка"] == expected


def test_a_transcript_longer_than_a_cell_is_cut_with_a_mark_inside_the_limit() -> None:
    """openpyxl режет молча на 32 767 — без пометки обрезанный разговор выглядел бы целым."""
    looped = "Клиент: алло, алло.\n" * 3000  # зацикливание распознавания, ~60 тыс. знаков
    (call,) = _table(_build([_row(transcript=looped)])["Звонки"])

    assert len(call["Расшифровка"]) == EXCEL_CELL_LIMIT
    assert call["Расшифровка"].endswith(TRUNCATED_MARK)
    assert transcript_counts([_row(transcript=looped), _row(transcript="Клиент: да.")]) == (2, 1)


def test_a_control_character_in_the_text_does_not_break_the_report() -> None:
    """Иначе `IllegalCharacterError` — 500 на весь отчёт и расшифровка в журнале ошибок."""
    (call,) = _table(
        _build([_row(transcript="Клиент: вот\x0bтак", summary="Абонент\x01 спросил.")])["Звонки"]
    )

    assert call["Расшифровка"] == "Клиент: воттак"
    assert call["Краткое содержание"] == "Абонент спросил."
    assert fit_text("а\x00б\x1fв") == "абв"


def test_a_transcript_that_looks_like_a_formula_stays_text() -> None:
    sheet = _build([_row(transcript="=1+1")])["Звонки"]
    header = [cell.value for cell in sheet[1]]

    cell = sheet.cell(row=2, column=header.index("Расшифровка") + 1)

    assert (cell.value, cell.data_type) == ("=1+1", "s")


def test_summaries_count_share_duration_sentiment_and_the_match() -> None:
    rows = [
        _row(operator_login="operator1", duration_s=100, sentiment="negative"),
        _row(operator_login="operator1", duration_s=300, sentiment="positive", ai_path=OPERATOR),
        _row(operator_login="operator2", duration_s=60, sentiment="neutral", ai_path=()),
        _row(operator_login=None, duration_s=10, sentiment=None),
    ]

    by_operator = _table(_build(rows)["По операторам"])

    assert by_operator[0] == {
        "Оператор": "operator1",
        "Звонков": 2,
        "Доля звонков, %": 50,
        "Средняя длительность, с": 200,
        "Негативных": 1,
        "Нейтральных": 0,
        "Позитивных": 1,
        "Совпадение ИИ с оператором, %": 90,
    }
    assert [row["Оператор"] for row in by_operator] == ["operator1", "operator2", "не указан"]
    assert by_operator[1]["Совпадение ИИ с оператором, %"] is None, "сравнивать не с чем"
    assert sum(row["Звонков"] for row in by_operator) == len(rows), (
        "звонок без оператора не потерян"
    )


def test_the_ai_classification_summary_names_the_whole_path() -> None:
    summary = _table(_build([_row(), _row(), _row(ai_path=())])["По классификации ИИ"])

    assert [(row["Классификация ИИ"], row["Звонков"]) for row in summary] == [
        (" › ".join(AI), 2),
        ("нет", 1),
    ]


def test_new_ai_wordings_are_marked_in_their_summary() -> None:
    rows = [
        _row(suggestion=("Физ.лицо", "ШПД", "Роутер"), suggestion_is_new=False),
        _row(suggestion=("Физ.лицо", "Мобильная связь", "Перенос номера"), suggestion_is_new=True),
    ]

    summary = {
        row["Тема по предложению ИИ помощника"]: row
        for row in _table(_build(rows)["По темам ИИ помощника"])
    }

    assert summary["Физ.лицо › Мобильная связь › Перенос номера"]["Новая формулировка"] == "да"
    assert summary["Физ.лицо › ШПД › Роутер"]["Новая формулировка"] in ("", None)


def test_text_that_looks_like_a_formula_stays_text() -> None:
    """Краткое содержание пишет модель по разговору: его содержимое нам не принадлежит."""
    sheet = _build([_row(summary='=HYPERLINK("http://example.invalid","x")')])["Звонки"]
    header = [cell.value for cell in sheet[1]]

    cell = sheet.cell(row=2, column=header.index("Краткое содержание") + 1)

    assert cell.value.startswith("=HYPERLINK")
    assert cell.data_type == "s"


def test_links_are_absolute_when_the_public_address_is_known() -> None:
    (call,) = _table(_build([_row()], base_url="https://speech.ic-group.ru/")["Звонки"])

    assert call["Ссылка"] == "https://speech.ic-group.ru/calls/b663725d-b3d8-4a96-8f0d-07c349ae6f5a"


def test_an_empty_period_is_still_a_file_with_headers() -> None:
    workbook = _build([])

    assert list(workbook["Звонки"].iter_rows(values_only=True))[0][0] == "Начало"
    assert summarize([], lambda r: r.operator_login, "не указан") == []
    match = _table(workbook["Совпадение с оператором"])
    assert match[-1]["Сравнено"] == 0


# --------------------------------------------------------------------------- #
# Причина обращения (S-32)
# --------------------------------------------------------------------------- #

K05 = CallReason(code="K05", label="Не работает домашний интернет")
K12 = CallReason(
    code="K12", label="Хочет узнать баланс, остаток, сумму или дату оплаты, задолженность"
)
K99 = CallReason(code="K99", label="Другое")
K07 = CallReason(code="K07", label="Не работает ТВ / приставка / не может войти в ТВ")


def _with_reason(code: str | None, **overrides: Any) -> ReportRow:
    labels = {reason.code: reason.label for reason in (K05, K12, K99, K07)}
    return _row(
        call_reason="Не работает домашний интернет." if code else None,
        call_reason_code=code,
        call_reason_label=labels.get(code or ""),
        **overrides,
    )


def test_the_reason_columns_stand_where_the_owners_sample_has_them() -> None:
    """AB и AC, как в образце: его формулы считают по колонке AC."""
    sheet = _build([_with_reason("K05")])["Звонки"]

    assert sheet["AB1"].value == "Причина обращения"
    assert sheet["AC1"].value == "Причина по справочнику"
    assert sheet["AB2"].value == "Не работает домашний интернет."
    assert sheet["AC2"].value == "K05 · Не работает домашний интернет"


def test_a_call_without_a_reason_is_not_given_a_code() -> None:
    """«Не разобран» — не K90: K90 значит «прочли, но не поняли»."""
    rows = [_with_reason(None, analysis_status="pending"), _with_reason(None)]

    calls = _table(_build(rows)["Звонки"])

    assert [call["Причина по справочнику"] for call in calls] == [
        "[не определена: звонок не разобран]",
        "[не определена]",
    ]
    assert [call["Причина обращения"] for call in calls] == [None, None]


def test_the_reasons_sheet_counts_every_call_and_lists_unused_active_codes() -> None:
    rows = [
        _with_reason("K05", duration_s=100),
        _with_reason("K05", duration_s=300),
        _with_reason("K12"),
        _with_reason("K07"),
        _with_reason(None, analysis_status="pending"),
    ]
    reasons = ((K05, True), (K12, True), (K99, True), (K07, False))

    summary = _table(_build(rows, reasons=reasons)["По причинам обращения"])

    names = [row["Причина обращения (справочник)"] for row in summary]
    assert names[0] == "K05 · Не работает домашний интернет"
    assert summary[0]["Звонков"] == 2
    assert summary[0]["Средняя длительность, с"] == 200
    assert "K07 · Не работает ТВ / приставка / не может войти в ТВ" in names, (
        "выключенный код, которым размечены звонки периода, не пропадает"
    )
    assert names[-1] == "K99 · Другое", "включённый код без звонков — строкой с нулём"
    assert summary[-1]["Звонков"] == 0
    assert "[не определена: звонок не разобран]" in names
    assert sum(row["Звонков"] for row in summary) == len(rows)


def test_the_reasons_sheet_holds_numbers_not_formulas() -> None:
    """Образец считал формулами `COUNTIF`; продукт формул не пишет (`_cell`)."""
    sheet = _build([_with_reason("K05")], reasons=((K05, True),))["По причинам обращения"]

    values = [cell.value for row in sheet.iter_rows(min_row=2) for cell in row]
    assert not any(isinstance(value, str) and value.startswith("=") for value in values)
