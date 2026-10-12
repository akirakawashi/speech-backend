"""Дашборд: сборка из суточных счётов и ручка, которая её отдаёт.

Сборка проверяется без базы: корзины по календарю, шесть тем и «Прочее», доля
негатива, половины периода, пик недели, подтемы. Ручка — на настоящем Postgres:
что темы складываются по веткам классификатора, что сутки считаются по часам
колл-центра и что в счёт идут только звонки NCC.

Тестам с базой нужен `SPEECH_POSTGRES_TEST_URL` с применённой схемой.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime, timedelta
from typing import Any

import httpx
import openpyxl
import pytest
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from src.application.dto.files import XLSX_MEDIA_TYPE
from src.di.http import create_app
from src.di.requests import session_dep
from src.domain.calls.contracts import SPEECH_ROLE_EDITOR, SPEECH_ROLE_READER
from src.infrastructure.persistence.classifier import NodePaths
from src.infrastructure.persistence.dashboard import (
    MAX_BUCKETS,
    OTHER_THEME,
    DayCount,
    auto_grain,
    bucket_spans,
    build,
    day_counts,
    resolve_grain,
)
from src.infrastructure.reports.dashboard import build_workbook
from src.presentation.http.auth import Principal, current_principal
from src.setting.config import Settings, get_settings
from tests.speech_db_support import speech_test_database_url

# --------------------------------------------------------------------------- #
# Сборка
# --------------------------------------------------------------------------- #

MON = date(2026, 9, 14)  # понедельник


def _row(
    day: date,
    theme: str,
    calls: int,
    *,
    negative: int = 0,
    rated: int | None = None,
    subtopic: str = "",
) -> DayCount:
    return DayCount(
        day=day,
        theme=theme.casefold(),
        theme_label=theme,
        subtopic=subtopic.casefold(),
        subtopic_label=subtopic,
        calls=calls,
        negative=negative,
        rated=calls if rated is None else rated,
    )


def test_the_grain_follows_the_length_of_the_period() -> None:
    assert [auto_grain(days) for days in (1, 10, 11, 120, 121, 400)] == [
        "day",
        "day",
        "week",
        "week",
        "month",
        "month",
    ]
    # Названная корзина действует, пока график её вмещает.
    year = (date(2026, 1, 1), date(2026, 12, 31))
    assert resolve_grain(MON, MON + timedelta(days=29), "day") == "day"
    assert resolve_grain(*year, "day") == "month"
    assert resolve_grain(*year, "week") == "week"
    assert len(bucket_spans(*year, "day")) > MAX_BUCKETS


def test_buckets_follow_the_calendar_and_are_cut_by_the_period() -> None:
    # Со среды по следующий вторник: две неполные недели.
    first, last = MON + timedelta(days=2), MON + timedelta(days=8)

    assert bucket_spans(first, last, "week") == [
        (first, MON + timedelta(days=6)),
        (MON + timedelta(days=7), last),
    ]
    assert bucket_spans(date(2026, 1, 30), date(2026, 3, 2), "month") == [
        (date(2026, 1, 30), date(2026, 1, 31)),
        (date(2026, 2, 1), date(2026, 2, 28)),
        (date(2026, 3, 1), date(2026, 3, 2)),
    ]
    assert len(bucket_spans(first, last, "day")) == 7


def test_bars_are_stacked_by_theme_and_a_day_without_calls_is_still_a_bar() -> None:
    rows = [
        _row(MON, "Не работает", 6, negative=3),
        _row(MON, "Тарифы", 2, negative=0),
        _row(MON + timedelta(days=2), "Не работает", 4, negative=1, rated=2),
    ]

    board = build(rows, first=MON, last=MON + timedelta(days=2))

    assert board.grain == "day"
    assert board.calls == 12
    assert [(theme.label, theme.calls, theme.share) for theme in board.themes] == [
        ("Не работает", 10, 0.8333),
        ("Тарифы", 2, 0.1667),
    ]
    assert [(bar.start, bar.calls, bar.by_theme) for bar in board.buckets] == [
        (MON, 8, (6, 2)),
        (MON + timedelta(days=1), 0, (0, 0)),
        (MON + timedelta(days=2), 4, (4, 0)),
    ]
    # Доля негатива — среди звонков с тональностью; где их нет, доли нет.
    assert [bar.negative_share for bar in board.buckets] == [0.375, None, 0.5]
    assert board.themes[0].negative_share == 0.5


def _themes(count: int) -> list[DayCount]:
    """`count` тем, каждая следующая реже предыдущей."""
    return [_row(MON, f"Тема {n}", 100 - n) for n in range(1, count + 1)]


def test_six_themes_stay_and_the_rest_fold_into_other() -> None:
    board = build(_themes(9), first=MON, last=MON)

    assert [theme.label for theme in board.themes] == [f"Тема {n}" for n in range(1, 7)] + [
        "Прочее"
    ]
    other = board.themes[-1]
    assert (other.key, other.calls) == (OTHER_THEME, 93 + 92 + 91)
    assert board.buckets[0].by_theme[-1] == other.calls
    assert sum(theme.calls for theme in board.themes) == board.calls
    # В выгрузку темы идут все и по отдельности.
    assert len(board.all_themes) == 9

    # Выбор «Прочего»: его числа, а подтемы — сложенные в него темы.
    chosen = build(_themes(9), first=MON, last=MON, theme=OTHER_THEME).selection
    assert (chosen.label, chosen.calls) == ("Прочее", other.calls)
    assert [(item.label, item.calls) for item in chosen.subtopics] == [
        ("Тема 7", 93),
        ("Тема 8", 92),
        ("Тема 9", 91),
    ]


def test_seven_themes_all_keep_their_names() -> None:
    """Прятать одну тему под «Прочее» незачем: место на графике у неё есть."""
    board = build(_themes(7), first=MON, last=MON)

    assert [theme.label for theme in board.themes] == [f"Тема {n}" for n in range(1, 8)]


def test_the_selection_counts_one_theme_or_all_of_them() -> None:
    last = MON + timedelta(days=13)
    rows = [
        _row(MON, "Не работает", 10, negative=5, subtopic="Wi-Fi"),
        _row(MON + timedelta(days=8), "Не работает", 20, negative=2, subtopic="Wi-Fi"),
        _row(MON + timedelta(days=8), "Не работает", 6, negative=0, subtopic="Кабель"),
        _row(MON + timedelta(days=9), "Не работает", 4, negative=0),
        _row(MON + timedelta(days=9), "Тарифы", 40, negative=4),
    ]

    everything = build(rows, first=MON, last=last).selection
    assert (everything.theme, everything.calls) == (None, 80)
    assert everything.negative_share == 0.1375
    # Первая неделя — 10 звонков, вторая — 70.
    assert everything.change == 6.0
    assert everything.peak is not None
    assert (everything.peak.start, everything.peak.end, everything.peak.calls) == (
        MON + timedelta(days=7),
        last,
        70,
    )
    assert everything.subtopics == ()

    chosen = build(rows, first=MON, last=last, theme="не работает").selection
    assert (chosen.label, chosen.calls, chosen.negative_share) == ("Не работает", 40, 0.175)
    assert chosen.change == 2.0
    # Подтемы — по убыванию, «Без уточнения» последней.
    assert [(item.label, item.calls, item.share) for item in chosen.subtopics] == [
        ("Wi-Fi", 30, 0.75),
        ("Кабель", 6, 0.15),
        ("Без уточнения", 4, 0.1),
    ]
    # Столбцы от выбора темы не зависят.
    assert build(rows, first=MON, last=last, theme="не работает").buckets == (
        build(rows, first=MON, last=last).buckets
    )


def test_a_theme_without_refinements_has_no_subtopics() -> None:
    chosen = build([_row(MON, "Тарифы", 5)], first=MON, last=MON, theme="тарифы").selection

    assert chosen.subtopics == ()


def test_a_theme_absent_from_the_period_is_zeros_not_an_error() -> None:
    """Ссылка с выбранной темой переживает смену периода."""
    board = build(
        [_row(MON, "Тарифы", 5)],
        first=MON,
        last=MON,
        theme="расторжение договора",
        known_labels={"расторжение договора": "Расторжение договора"},
    )

    assert (board.selection.label, board.selection.calls) == ("Расторжение договора", 0)
    assert (board.selection.change, board.selection.peak) == (None, None)


def test_the_halves_of_an_odd_period_leave_the_middle_day_out() -> None:
    rows = [
        _row(MON, "Тема", 10),
        _row(MON + timedelta(days=1), "Тема", 500),
        _row(MON + timedelta(days=2), "Тема", 15),
    ]

    assert build(rows, first=MON, last=MON + timedelta(days=2)).selection.change == 0.5
    # Один день пополам не делится; пустую первую половину сравнивать не с чем.
    assert build(rows, first=MON, last=MON).selection.change is None
    late = [_row(MON + timedelta(days=1), "Тема", 5)]
    assert build(late, first=MON, last=MON + timedelta(days=1)).selection.change is None


def test_calls_outside_the_period_are_left_out() -> None:
    rows = [_row(MON - timedelta(days=1), "Тема", 7), _row(MON, "Тема", 3)]

    assert build(rows, first=MON, last=MON).calls == 3


def _tree(*paths: tuple[str, ...]) -> tuple[NodePaths, dict[tuple[str, ...], uuid.UUID]]:
    """Дерево из путей: узел на каждый префикс пути."""
    ids: dict[tuple[str, ...], uuid.UUID] = {}
    rows = []
    for path in paths:
        for depth in range(1, len(path) + 1):
            prefix = path[:depth]
            if prefix not in ids:
                ids[prefix] = uuid.uuid4()
                rows.append(
                    {
                        "id": ids[prefix],
                        "parent_id": ids.get(prefix[:-1]),
                        "level": depth,
                        "label": prefix[-1],
                        "sort_order": 0,
                        "is_active": True,
                    }
                )
    return NodePaths(rows), ids


def test_a_theme_is_one_name_across_branches_and_an_unnamed_level_is_not_a_name() -> None:
    home = ("Физ.лицо", "ШПД", "Консультация", "Не работает")
    paths, ids = _tree(
        (*home, "Wi-Fi"),
        # Прочерк пятым уровнем ставят операторы там, где уточнения нет.
        (*home, "-"),
        # Та же тема в другой ветке и в другом написании — та же тема.
        ("В2В", "Мобильная связь", "Консультация", "не работает"),
        # Уровень темы без названия — не тема.
        ("Физ.лицо", "ШПД", "Недовольство/негатив", "-", "Что-то"),
    )

    def row(*path: str) -> dict[str, Any]:
        return {"day": MON, "node_id": ids[path], "calls": 1, "negative": 0, "rated": 1}

    counted = day_counts(
        [
            row(*home, "Wi-Fi"),
            row(*home, "-"),
            row(*home),
            row("В2В", "Мобильная связь", "Консультация", "не работает"),
            row("Физ.лицо", "ШПД", "Недовольство/негатив", "-", "Что-то"),
            # Путь не дошёл до уровня темы.
            row("Физ.лицо", "ШПД", "Консультация"),
            # Узла нет в дереве: его удалили, а звонок остался.
            {"day": MON, "node_id": uuid.uuid4(), "calls": 9, "negative": 0, "rated": 9},
        ],
        paths,
    )

    assert {item.theme for item in counted} == {"не работает"}
    assert sorted(item.subtopic_label for item in counted) == ["", "", "", "Wi-Fi"]
    chosen = build(counted, first=MON, last=MON, theme="не работает").selection
    assert [(item.label, item.calls) for item in chosen.subtopics] == [
        ("Wi-Fi", 1),
        ("Без уточнения", 3),
    ]


def test_a_theme_named_like_a_formula_stays_text_in_the_file() -> None:
    """Названия тем — из классификатора, а его правят люди: формулой они не становятся."""
    rows = [
        _row(MON, "=1+1", 3, subtopic="=SUM(A1)"),
        _row(MON, "Тарифы", 2),
    ]

    book = openpyxl.load_workbook(io.BytesIO(build_workbook(build(rows, first=MON, last=MON))))

    for sheet in book:
        for line in sheet.iter_rows():
            for cell in line:
                assert cell.data_type != "f", f"{sheet.title}!{cell.coordinate}"
    # И в шапке листа тоже: тема там — заголовок столбца.
    assert book["По времени"]["C1"].value == "=1+1"
    assert book["Темы"]["A2"].value == "=1+1"
    assert book["Подтемы"]["B2"].value == "=SUM(A1)"


# --------------------------------------------------------------------------- #
# Ручка
# --------------------------------------------------------------------------- #

_TEST_DATABASE_URL = speech_test_database_url()
needs_db = pytest.mark.skipif(not _TEST_DATABASE_URL, reason="SPEECH_POSTGRES_TEST_URL is required")

Sessions = async_sessionmaker[AsyncSession]
START = datetime(2026, 9, 14, tzinfo=UTC)
PERIOD = {
    # Так границы присылает экран: полночь по Москве и полночь следующих суток.
    "from": "2026-09-13T21:00:00Z",
    "to": "2026-09-16T21:00:00Z",
}
TABLES = ("speech_audit_events", "speech_calls", "speech_classifier_nodes")


async def _clean(sessions: Sessions) -> None:
    async with sessions() as session, session.begin():
        for table in TABLES:
            await session.execute(text(f"DELETE FROM {table}"))  # noqa: S608


@pytest.fixture()
async def sessions() -> AsyncIterator[Sessions]:
    assert _TEST_DATABASE_URL is not None
    engine = create_async_engine(_TEST_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    await _clean(factory)
    yield factory
    await _clean(factory)
    await engine.dispose()


async def _path(sessions: Sessions, *labels: str) -> uuid.UUID:
    """Завести путь классификатора и вернуть его последний узел."""
    parent: uuid.UUID | None = None
    async with sessions() as session, session.begin():
        for level, label in enumerate(labels, start=1):
            found = await session.scalar(
                text(
                    "SELECT id FROM speech_classifier_nodes WHERE level = :level "
                    "AND label = :label AND parent_id IS NOT DISTINCT FROM :parent"
                ),
                {"level": level, "label": label, "parent": parent},
            )
            if found is None:
                found = uuid.uuid4()
                await session.execute(
                    text(
                        "INSERT INTO speech_classifier_nodes (id, parent_id, level, label, key, "
                        "source, sort_order, is_active, created_at, updated_at) "
                        "VALUES (:id, :parent, :level, :label, :key, 'file', 0, true, now(), now())"
                    ),
                    {
                        "id": found,
                        "parent": parent,
                        "level": level,
                        "label": label,
                        "key": label.casefold(),
                    },
                )
            parent = found
    assert parent is not None
    return parent


async def _calls(
    sessions: Sessions,
    node: uuid.UUID | None,
    count: int,
    *,
    started: datetime,
    sentiment: str | None = "neutral",
    source: str = "ncc",
) -> None:
    async with sessions() as session, session.begin():
        for _ in range(count):
            await session.execute(
                text(
                    "INSERT INTO speech_calls (id, session_id, source, direction, started_at, "
                    "ended_at, duration_s, audio_status, processing_status, analysis_status, "
                    "segment_source, product_source, source_snapshot, ai_node_id, sentiment, "
                    "first_seen_at, created_at, updated_at) "
                    "VALUES (gen_random_uuid(), :sid, :source, 'inbound', :s, :s, 60, 'stored', "
                    "'done', 'done', 'unknown', 'unknown', '{}', :node, :sentiment, :s, :s, :s)"
                ),
                {
                    "sid": f"node_dash_{uuid.uuid4().hex}",
                    "source": source,
                    "s": started,
                    "node": node,
                    "sentiment": sentiment,
                },
            )


def _client(sessions: Sessions, *roles: str) -> httpx.AsyncClient:
    settings = Settings(
        _env_file=None,  # type: ignore[call-arg]
        app_env="local",
        speech_auth_mode="open",
        speech_ms_database_url=_TEST_DATABASE_URL or "",
    )

    async def _override() -> AsyncIterator[AsyncSession]:
        async with sessions() as session:
            yield session

    app = create_app(settings)
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[session_dep] = _override
    if roles:
        app.dependency_overrides[current_principal] = lambda: Principal(
            sub="someone",
            username="someone",
            email=None,
            display_name=None,
            roles=roles,
            auth_mode="keycloak",
        )
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://speech")


@pytest.fixture()
async def seeded(sessions: Sessions) -> Sessions:
    """Три дня звонков: одна тема в двух ветках дерева, вторая тема и шум вокруг."""
    wifi = await _path(sessions, "Физ.лицо", "ШПД", "Консультация", "Не работает", "Wi-Fi")
    # Та же тема у другого сегмента и продукта — другой узел, та же тема.
    mobile = await _path(sessions, "В2В", "Мобильная связь", "Консультация", "Не работает")
    tariffs = await _path(sessions, "Физ.лицо", "ШПД", "Консультация", "Тарифы")
    shallow = await _path(sessions, "Физ.лицо", "ШПД", "Консультация")

    noon = START + timedelta(hours=9)
    await _calls(sessions, wifi, 3, started=noon, sentiment="negative")
    await _calls(sessions, mobile, 1, started=noon)
    await _calls(sessions, tariffs, 2, started=noon + timedelta(days=2))
    # 23:30 по Москве 15-го — это ещё 15-е, хотя по UTC уже 20:30 того же дня.
    await _calls(sessions, tariffs, 1, started=datetime(2026, 9, 15, 20, 30, tzinfo=UTC))
    # Без тональности: в счёт обращений идёт, в долю негатива — нет.
    await _calls(sessions, wifi, 1, started=noon, sentiment=None)
    # Не входят: звонок Телфина, звонок без темы, путь до темы не дошёл, вне периода.
    await _calls(sessions, wifi, 5, started=noon, source="telphin")
    await _calls(sessions, None, 4, started=noon)
    await _calls(sessions, shallow, 2, started=noon)
    await _calls(sessions, wifi, 9, started=START - timedelta(days=5))
    return sessions


@needs_db
async def test_the_dashboard_counts_ncc_calls_by_theme_and_day(seeded: Sessions) -> None:
    async with _client(seeded) as client:
        response = await client.get("/api/v1/dashboard", params=PERIOD)

    assert response.status_code == 200
    body = response.json()
    assert (body["first_day"], body["last_day"], body["grain"]) == (
        "2026-09-14",
        "2026-09-16",
        "day",
    )
    assert body["calls"] == 8
    assert [(theme["label"], theme["calls"]) for theme in body["themes"]] == [
        ("Не работает", 5),
        ("Тарифы", 3),
    ]
    # Три негативных из четырёх звонков темы с тональностью.
    assert body["themes"][0]["negative_share"] == 0.75
    assert [(bar["start"], bar["calls"], bar["by_theme"]) for bar in body["buckets"]] == [
        ("2026-09-14", 5, [5, 0]),
        ("2026-09-15", 1, [0, 1]),
        ("2026-09-16", 2, [0, 2]),
    ]
    assert body["selection"]["theme"] is None
    assert body["selection"]["calls"] == 8


@needs_db
async def test_a_chosen_theme_comes_with_its_subtopics(seeded: Sessions) -> None:
    async with _client(seeded) as client:
        chosen = (
            await client.get("/api/v1/dashboard", params={**PERIOD, "theme": "не работает"})
        ).json()["selection"]
        weeks = (await client.get("/api/v1/dashboard", params={**PERIOD, "grain": "week"})).json()
        empty = await client.get(
            "/api/v1/dashboard",
            params={"from": "2026-09-16T21:00:00Z", "to": "2026-09-13T21:00:00Z"},
        )

    assert (chosen["label"], chosen["calls"], chosen["negative_share"]) == ("Не работает", 5, 0.75)
    assert [(item["label"], item["calls"]) for item in chosen["subtopics"]] == [
        ("Wi-Fi", 4),
        ("Без уточнения", 1),
    ]
    assert chosen["peak"] == {"start": "2026-09-14", "end": "2026-09-16", "calls": 5}
    assert (weeks["grain"], len(weeks["buckets"])) == ("week", 1)
    assert (empty.status_code, empty.json()["detail"]) == (422, "period_is_empty")


@needs_db
async def test_the_file_carries_the_same_numbers_and_needs_the_export_right(
    seeded: Sessions,
) -> None:
    async with _client(seeded, SPEECH_ROLE_READER) as reader:
        seen = await reader.get("/api/v1/dashboard", params=PERIOD)
        refused = await reader.get("/api/v1/dashboard.xlsx", params=PERIOD)
    async with _client(seeded, SPEECH_ROLE_EDITOR) as editor:
        response = await editor.get("/api/v1/dashboard.xlsx", params=PERIOD)

    # Читатель дашборд видит, а файл наружу не выносит.
    assert (seen.status_code, refused.status_code) == (200, 403)
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="dashboard-2026-09-14_2026-09-16.xlsx"'
    )
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["По времени", "Темы", "Подтемы"]

    def rows(sheet: str) -> list[tuple[Any, ...]]:
        return list(book[sheet].iter_rows(values_only=True))

    assert rows("По времени")[0] == (
        "День, с",
        "по",
        "Не работает",
        "Тарифы",
        "Всего",
        "Доля негатива",
    )
    assert rows("По времени")[1] == (
        datetime(2026, 9, 14),
        datetime(2026, 9, 14),
        5,
        0,
        5,
        0.75,
    )
    assert rows("Темы")[1:] == [("Не работает", 5, 0.625, 0.75), ("Тарифы", 3, 0.375, 0)]
    assert rows("Подтемы")[1:] == [
        ("Не работает", "Wi-Fi", 4, 0.8),
        ("Не работает", "Без уточнения", 1, 0.2),
    ]

    async with seeded() as session:
        events = (
            await session.execute(
                text(
                    "SELECT actor_sub, payload FROM speech_audit_events "
                    "WHERE event_type = 'report.dashboard'"
                )
            )
        ).all()
    assert [event[0] for event in events] == ["someone"]
