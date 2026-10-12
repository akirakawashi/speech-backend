"""Файл «Сверка оценок»: сборка книги и ручка, которая её отдаёт.

Сборка проверяется без базы — по книге, прочитанной обратно. Ручка — на
настоящем Postgres: какие звонки попадают в выборку, что в файл идёт текст
расшифровки как есть, что выгрузка пишется в журнал и требует права `export`.

Тестам с базой нужен `SPEECH_POSTGRES_TEST_URL` с применённой схемой.
"""

from __future__ import annotations

import io
import json
import re
import uuid
from collections.abc import AsyncIterator, Awaitable, Callable
from datetime import UTC, datetime, timedelta
from typing import Any

import httpx
import openpyxl
import pytest
from sqlalchemy import text

from src.application import authorization
from src.application.dto.files import XLSX_MEDIA_TYPE
from src.application.quality.batch import evaluate_calls
from src.domain.calls.contracts import SPEECH_ROLE_EDITOR, SPEECH_ROLE_READER
from src.domain.quality.script import CURRENT_SCRIPT
from src.infrastructure.persistence.evaluation import PostgresEvaluationRepository
from src.infrastructure.persistence.quality_read_operations import _review_filename
from src.infrastructure.reports.quality import NOT_NEEDED, ReviewCall, ReviewLine, build_workbook
from src.presentation.http import auth as auth
from src.presentation.http.auth import PERM_QUALITY, Principal
from tests.speech_db_support import speech_test_database_url
from tests.test_speech_quality_db import (  # noqa: F401 - фикстуры
    END,
    START,
    Sessions,
    _answer,
    _api_client,
    _call,
    _FakeModel,
    _operator,
    _refs,
    _retranscribe,
    api,
    sessions,
)

SCRIPT = CURRENT_SCRIPT
STAGES = len(SCRIPT.stages)

needs_db = pytest.mark.skipif(
    not speech_test_database_url(), reason="SPEECH_POSTGRES_TEST_URL is required"
)

# --------------------------------------------------------------------------- #
# Сборка книги
# --------------------------------------------------------------------------- #

LINES = (
    ReviewLine(seq=0, speaker="operator", start_ms=0, text="Добрый день, компания Волна"),
    ReviewLine(seq=1, speaker="client", start_ms=4_000, text="Мне это не интересно"),
    ReviewLine(seq=2, speaker="unknown", start_ms=65_000, text="Угу"),
)


def _verdicts(*, failed: tuple[str, ...] = (), comment: str = "") -> list[dict[str, Any]]:
    return [
        {
            "key": stage.key,
            "applicable": stage.key != "busy",
            "passed": stage.key not in failed,
            "evidence": [0, 1] if stage.key in failed else [0],
            "comment": comment if stage.key in failed else "",
        }
        for stage in SCRIPT.stages
    ]


def _review_call(
    *,
    started: datetime = datetime(2026, 9, 14, 9, 30, tzinfo=UTC),
    operator: str = "Анна Петрова",
    stages: list[dict[str, Any]] | None = None,
    lines: tuple[ReviewLine, ...] = LINES,
) -> ReviewCall:
    return ReviewCall(
        call_id=uuid.uuid4(),
        started_at=started,
        duration_s=205,
        operator=operator,
        stages=_verdicts() if stages is None else stages,
        lines=lines,
    )


def _book(content: bytes) -> Any:
    return openpyxl.load_workbook(io.BytesIO(content))


def _build(*calls: ReviewCall, base_url: str = "https://speech.example") -> Any:
    return _book(build_workbook(calls, script=SCRIPT, timezone="Europe/Moscow", base_url=base_url))


def _rows(sheet: Any) -> list[dict[str, Any]]:
    """Строки листа словарями по заголовку."""
    names = [cell.value for cell in sheet[1]]
    return [
        dict(zip(names, values, strict=True))
        for values in sheet.iter_rows(min_row=2, values_only=True)
    ]


MARK = "Оценка модели"
STAGE = "Этап скрипта"
QUOTES = "Реплики, на которые сослалась модель"
CONTROLLER = "Оценка контролёра (1 / 0 / не понадобился)"


def test_each_stage_is_a_row_with_the_mark_the_reason_and_the_quoted_replies() -> None:
    call = _review_call(stages=_verdicts(failed=("passport",), comment="Про паспорт не сказал."))

    book = _build(call)

    assert book.sheetnames == ["Сверка", "Разговоры", "Этапы"]
    rows = _rows(book["Сверка"])
    # Этапы — в порядке скрипта и с его номерами.
    assert [row[STAGE] for row in rows] == [
        f"{number}. {stage.label}" for number, stage in enumerate(SCRIPT.stages, start=1)
    ]
    by_key = dict(zip((stage.key for stage in SCRIPT.stages), rows, strict=True))
    passport, greeting, busy = by_key["passport"], by_key["greeting"], by_key["busy"]
    assert (passport[MARK], passport["Обязательный"]) == (0, "да")
    assert passport["Пояснение модели"] == "Про паспорт не сказал."
    assert passport[QUOTES] == (
        "00:00 Оператор: Добрый день, компания Волна\n00:04 Клиент: Мне это не интересно"
    )
    assert (greeting[MARK], greeting["Обязательный"]) == (1, None)
    # Не понадобившийся этап — не единица: контролёр ставит его третьим значением.
    assert busy[MARK] == NOT_NEEDED
    # Звонок описан в каждой строке: её можно фильтровать и сортировать отдельно.
    assert {row["Сотрудник"] for row in rows} == {"Анна Петрова"}
    assert {row["Длительность"] for row in rows} == {"03:25"}
    # 09:30 UTC — это 12:30 в колл-центре.
    assert {row["Звонок"] for row in rows} == {datetime(2026, 9, 14, 12, 30)}
    # Колонки контролёра пусты: их заполняет он.
    assert {(row[CONTROLLER], row["Комментарий контролёра"]) for row in rows} == {(None, None)}


def test_the_whole_conversation_is_on_its_own_sheet() -> None:
    """«Этого не сказали» по цитатам не проверить — нужен весь разговор."""
    talks = _rows(_build(_review_call())["Разговоры"])

    assert [(row["Время"], row["Кто говорит"], row["Реплика"]) for row in talks] == [
        ("00:00", "Оператор", "Добрый день, компания Волна"),
        ("00:04", "Клиент", "Мне это не интересно"),
        ("01:05", "Не определено", "Угу"),
    ]
    assert {row["№"] for row in talks} == {1}


def test_calls_are_numbered_by_time_and_alike_on_both_sheets() -> None:
    late = _review_call(started=datetime(2026, 9, 15, 8, tzinfo=UTC), operator="Олег")
    early = _review_call(started=datetime(2026, 9, 14, 8, tzinfo=UTC), operator="")

    book = _build(late, early)

    def numbers(sheet: str) -> dict[int, str]:
        return {row["№"]: row["Сотрудник"] for row in _rows(book[sheet])}

    # Выборка случайна, а читают файл по времени. Звонок без сотрудника назван так.
    assert numbers("Сверка") == numbers("Разговоры") == {1: "Без сотрудника", 2: "Олег"}


def test_text_from_the_call_never_becomes_a_formula() -> None:
    """Реплика и пояснение — чужой текст: формула из них выполнилась бы у контролёра."""
    lines = (ReviewLine(seq=0, speaker="client", start_ms=0, text='=HYPERLINK("http://x")'),)
    stages = _verdicts(failed=("passport",), comment="=1+1")
    book = _build(_review_call(stages=stages, lines=lines))

    for sheet in book:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"{sheet.title}!{cell.coordinate}"
    passport = next(row for row in _rows(book["Сверка"]) if row[MARK] == 0)
    assert passport["Пояснение модели"] == "=1+1"
    assert _rows(book["Разговоры"])[0]["Реплика"] == '=HYPERLINK("http://x")'


def test_a_control_character_does_not_break_the_file() -> None:
    lines = (ReviewLine(seq=0, speaker="client", start_ms=0, text="Алло\x0bслышно"),)

    talks = _rows(_build(_review_call(lines=lines))["Разговоры"])

    assert talks[0]["Реплика"] == "Аллослышно"


def test_the_link_opens_the_review_window_of_the_call() -> None:
    call = _review_call()

    linked = _build(call, base_url="https://speech.example/")["Сверка"]["E2"]
    plain = _build(call, base_url="")["Сверка"]["E2"]

    assert linked.value == "Открыть"
    assert linked.hyperlink.target == f"https://speech.example/quality?call={call.call_id}"
    # Без адреса сервиса ссылку открыть нечем: путь остаётся текстом.
    assert (plain.value, plain.hyperlink) == (f"/quality?call={call.call_id}", None)


def test_a_reply_number_outside_the_call_is_skipped_not_fatal() -> None:
    stages = [{**item, "evidence": [99, 1, "x"]} for item in _verdicts()]

    rows = _rows(_build(_review_call(stages=stages))["Сверка"])

    assert {row[QUOTES] for row in rows} == {"00:04 Клиент: Мне это не интересно"}


def test_the_stage_reference_says_what_the_model_was_asked_to_check() -> None:
    reference = _rows(_build(_review_call())["Этапы"])

    assert [row["Что требуется"] for row in reference] == [
        stage.requirement for stage in SCRIPT.stages
    ]
    passport = next(row for row in reference if "паспорт" in row[STAGE])
    assert passport["Обязательный"] == "да"
    assert passport["Когда этап нужен"] == "Оформлена заявка на подключение."
    assert reference[0]["Когда этап нужен"] == "В любом звонке по скрипту."


def test_the_file_name_carries_the_days_as_the_call_centre_sees_them() -> None:
    # Так границы присылает экран: начало суток в поясе показа, верхняя —
    # начало следующих суток.
    start = datetime(2026, 9, 13, 21, tzinfo=UTC)
    end = datetime(2026, 9, 20, 21, tzinfo=UTC)

    assert _review_filename(start, end, "Europe/Moscow") == (
        "quality-review-2026-09-14_2026-09-20.xlsx"
    )
    assert _review_filename(None, None, "Europe/Moscow") == "quality-review.xlsx"
    assert _review_filename(start, None, "Europe/Moscow") == "quality-review.xlsx"


# --------------------------------------------------------------------------- #
# Ручка
# --------------------------------------------------------------------------- #

URL = "/api/v1/quality/review.xlsx"
PERIOD = {"from": START.isoformat(), "to": END.isoformat()}
EVENT = "report.quality_review"

Journal = Callable[[], Awaitable[list[dict[str, Any]]]]


@pytest.fixture()
async def journal(sessions: Sessions) -> AsyncIterator[Journal]:  # noqa: F811
    """События выгрузки в журнале. Записи теста после него убираются."""

    async def drop() -> None:
        async with sessions() as session, session.begin():
            await session.execute(
                text("DELETE FROM speech_audit_events WHERE event_type = :event"),
                {"event": EVENT},
            )

    async def read() -> list[dict[str, Any]]:
        async with sessions() as session:
            rows = await session.execute(
                text(
                    "SELECT actor_sub, payload FROM speech_audit_events "
                    "WHERE event_type = :event ORDER BY occurred_at"
                ),
                {"event": EVENT},
            )
            return [
                {
                    "actor": row["actor_sub"],
                    "payload": (
                        row["payload"]
                        if isinstance(row["payload"], dict)
                        else json.loads(row["payload"])
                    ),
                }
                for row in rows.mappings()
            ]

    await drop()
    yield read
    await drop()


def _call_ids(sheet: Any) -> set[str]:
    """Звонки листа «Сверка» — по ссылке на разбор."""
    found = set()
    for row in _rows(sheet):
        match = re.search(r"call=([0-9a-f-]{36})", str(row["Запись и разбор"]))
        assert match is not None
        found.add(match.group(1))
    return found


def _all_text(book: Any) -> str:
    return "\n".join(
        str(value)
        for sheet in book
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    )


@needs_db
async def test_the_file_is_a_random_sample_of_the_calls_scored_by_the_script(
    sessions: Sessions,  # noqa: F811
    api: httpx.AsyncClient,  # noqa: F811
    journal: Journal,
) -> None:
    anna = await _operator(sessions, "anna")
    scored = [
        await _call(
            sessions, f"scored{n}", started=START + timedelta(hours=9 + n), operator_id=anna
        )
        for n in range(3)
    ]
    inbound = await _call(sessions, "inbound", started=START + timedelta(hours=13))
    broken = await _call(sessions, "broken", started=START + timedelta(hours=14))
    await _call(sessions, "pending", started=START + timedelta(hours=15))
    for call_id, answer in (
        *((call_id, _answer(failed=["passport"])) for call_id in scored),
        (inbound, _answer(applies=False)),
        (broken, "not JSON"),
    ):
        await evaluate_calls(
            PostgresEvaluationRepository(sessions),
            _FakeModel(answer),
            await _refs(sessions, call_id),
        )

    # Без лимита в файл идут все, кого можно сверять: ровно три оценённых по
    # скрипту — ни звонка вне скрипта, ни неудавшейся оценки, ни ожидающего.
    everyone = await api.get(URL, params=PERIOD)
    assert everyone.status_code == 200
    assert _call_ids(_book(everyone.content)["Сверка"]) == {str(call_id) for call_id in scored}

    response = await api.get(URL, params={**PERIOD, "limit": 2})

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"].startswith(
        'attachment; filename="quality-review-2026-09-14_'
    )
    book = _book(response.content)
    check = _rows(book["Сверка"])
    # С лимитом — случайные два из тех же трёх.
    picked = _call_ids(book["Сверка"])
    assert len(picked) == 2 and picked <= {str(call_id) for call_id in scored}
    assert len(check) == 2 * STAGES
    assert {row["Сотрудник"] for row in check} == {"Сотрудник anna"}
    passport = [row for row in check if "паспорт" in row[STAGE]]
    assert [row[MARK] for row in passport] == [0, 0]
    assert passport[0][QUOTES] == "00:00 Оператор: Меня зовут Анна, компания Волна"
    assert {row[MARK] for row in check if "паспорт" not in row[STAGE]} == {1, NOT_NEEDED}

    # Текст — как есть, тот же, что видела модель.
    everything = _all_text(book)
    assert "79781234567" in everything and "Анна" in everything
    assert len(_rows(book["Разговоры"])) == 2 * 2

    # Файл ушёл из системы — это видно в журнале.
    events = await journal()
    assert [event["payload"]["rows"] for event in events] == [3, 2]
    assert events[0]["payload"]["script_version"] == SCRIPT.version


@needs_db
async def test_the_file_follows_the_scope_of_the_screen(
    sessions: Sessions,  # noqa: F811
    api: httpx.AsyncClient,  # noqa: F811
    journal: Journal,
) -> None:
    anna = await _operator(sessions, "anna")
    oleg = await _operator(sessions, "oleg")
    hers = await _call(sessions, "hers", started=START + timedelta(hours=9), operator_id=anna)
    his = await _call(sessions, "his", started=START + timedelta(hours=10), operator_id=oleg)
    nobody = await _call(sessions, "nobody", started=START + timedelta(hours=11))
    for call_id in (hers, his, nobody):
        await evaluate_calls(
            PostgresEvaluationRepository(sessions),
            _FakeModel(_answer()),
            await _refs(sessions, call_id),
        )

    async def picked(**params: str) -> set[str]:
        response = await api.get(URL, params={**PERIOD, **params})
        assert response.status_code == 200
        return _call_ids(_book(response.content)["Сверка"])

    assert await picked() == {str(hers), str(his), str(nobody)}
    assert await picked(operator_id=str(oleg)) == {str(his)}
    assert await picked(unassigned="true") == {str(nobody)}

    # Звонок расшифровали заново: прежняя оценка устарела, и сверять по ней нечего.
    await _retranscribe(sessions, his)
    assert await picked() == {str(hers), str(nobody)}

    # В охвате нет оценённых звонков — отказ с причиной, а не пустой файл, и в
    # журнале такой запрос выгрузкой не числится.
    before = len(await journal())
    earlier = {"from": (START - timedelta(days=7)).isoformat(), "to": START.isoformat()}
    empty = await api.get(URL, params=earlier)
    assert (empty.status_code, empty.json()["detail"]) == (404, "nothing_to_review")
    assert len(await journal()) == before


@needs_db
async def test_seeing_the_scores_is_not_enough_to_take_the_file_out(
    sessions: Sessions,  # noqa: F811
    journal: Journal,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Файл уходит из системы: кроме `quality` нужно право `export`."""
    call_id = await _call(sessions, "scored")
    await evaluate_calls(
        PostgresEvaluationRepository(sessions),
        _FakeModel(_answer()),
        await _refs(sessions, call_id),
    )
    # Роли «оценки без выгрузки» сегодня нет — заводим её на время теста: проверка
    # должна держаться и тогда, когда роли разделят.
    monkeypatch.setitem(authorization._PERMISSIONS, "quality_only", frozenset({PERM_QUALITY}))

    async def status_for(*roles: str) -> int:
        principal = Principal(
            sub="someone",
            username="someone",
            email=None,
            display_name=None,
            roles=roles,
            auth_mode="keycloak",
        )
        async with _api_client(sessions, principal=principal) as client:
            return (await client.get(URL, params=PERIOD)).status_code

    assert await status_for(SPEECH_ROLE_READER) == 403
    assert await status_for("quality_only") == 403
    assert await journal() == []
    assert await status_for(SPEECH_ROLE_EDITOR) == 200
    assert [event["actor"] for event in await journal()] == ["someone"]
